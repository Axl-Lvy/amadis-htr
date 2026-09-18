"""The human-assigned reference for where each Trésor piece sits in the baseline.

Read from the reference workbook's column B, which is prose written by an editor
rather than a clean field. Square brackets mark a conjecture and a question mark
marks doubt, so both survive into a confidence column instead of being flattened
away.

This is the only ground truth E5 may use. The `Alignment` table in amadis cannot
serve, because a human's pick and the matcher's own past output are
indistinguishable there.
"""

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl

CONFIDENCES: tuple[str, ...] = ("certain", "conjecture", "uncertain", "none")

_LOCATION = re.compile(
    r"livre\s+([0-9]+)(?:\s*,\s*chap\.?\s*([0-9]+))?", re.IGNORECASE
)
_QUESTION = re.compile(r"\(?\?+\)?")
_PIECE = re.compile(r"^\s*([0-9]+)\s*\.")
_SHEET = "Feuil1"


@dataclass(frozen=True)
class Reference:
    piece: int
    livre: int | None
    chapter: int | None
    confidence: str
    raw: str


def classify(raw: object) -> tuple[str, int | None, int | None]:
    """Turn one column B cell into (confidence, livre, chapter)."""
    if raw is None:
        return ("none", None, None)

    text = str(raw).strip()
    if not text or text.upper() == "N/A":
        return ("none", None, None)

    confidence = "certain"
    if text.startswith("[") and text.endswith("]"):
        confidence = "conjecture"
        text = text[1:-1].strip()
    if "?" in text:
        # A bracketed conjecture that also carries a question mark stays a
        # conjecture: the editor's bracket is the stronger signal.
        if confidence != "conjecture":
            confidence = "uncertain"
        text = _QUESTION.sub("", text).strip()

    text = text.rstrip(". ").strip()
    match = _LOCATION.search(text)
    if match is None:
        return ("none", None, None)

    livre = int(match.group(1))
    chapter = int(match.group(2)) if match.group(2) else None
    return (confidence, livre, chapter)


def read_workbook(path: str | Path) -> list[Reference]:
    """Read every row of the reference workbook, in piece order."""
    sheet = openpyxl.load_workbook(path, data_only=True)[_SHEET]
    references: list[Reference] = []
    for row in range(1, sheet.max_row + 1):
        title = str(sheet.cell(row, 1).value or "")
        piece_match = _PIECE.match(title)
        if piece_match is None:
            # Trailing blank rows end the table. Anything else in column A is a
            # row the editor added in a shape this parser does not understand,
            # and silently skipping it would lose a reference.
            if not title.strip():
                break
            raise ValueError(
                f"row {row}: column A does not start with a number: {title!r}"
            )
        cell = sheet.cell(row, 2).value
        confidence, livre, chapter = classify(cell)
        references.append(
            Reference(
                piece=int(piece_match.group(1)),
                livre=livre,
                chapter=chapter,
                confidence=confidence,
                raw="" if cell is None else str(cell).strip(),
            )
        )
    return references


def read_titles(path: str | Path) -> dict[int, str]:
    """Read column A as piece number to title, with the number stripped off.

    Gate G1 matches these titles against the Juxtalinéaire docx. They stay in
    this module because the workbook's sheet name and row shape are declared
    here and nowhere else.
    """
    sheet = openpyxl.load_workbook(path, data_only=True)[_SHEET]
    titles: dict[int, str] = {}
    for row in range(1, sheet.max_row + 1):
        cell = str(sheet.cell(row, 1).value or "")
        match = _PIECE.match(cell)
        if match is None:
            if not cell.strip():
                break
            raise ValueError(
                f"row {row}: column A does not start with a number: {cell!r}"
            )
        titles[int(match.group(1))] = " ".join(cell[match.end() :].split())
    return titles


def write_csv(references: Iterable[Reference], path: str | Path) -> None:
    """Write the reference table, one row per piece."""
    rows: Sequence[Reference] = list(references)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["piece", "livre", "chapter", "confidence", "raw"])
        for ref in rows:
            writer.writerow(
                [
                    ref.piece,
                    "" if ref.livre is None else ref.livre,
                    "" if ref.chapter is None else ref.chapter,
                    ref.confidence,
                    ref.raw,
                ]
            )
